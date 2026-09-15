from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_blocks(parent):
    if isinstance(parent, DocumentObject):
        element = parent.element.body
    else:
        raise TypeError(type(parent))
    for child in element.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, parent)
        elif child.tag.endswith("}tbl"):
            yield Table(child, parent)


def extract(path: Path) -> dict:
    doc = Document(path)
    blocks = []
    for block in iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                blocks.append({"type": "paragraph", "style": block.style.name, "text": text})
        else:
            rows = []
            for row in block.rows:
                rows.append([cell.text.strip() for cell in row.cells])
            blocks.append({"type": "table", "rows": rows})

    headers = []
    footers = []
    for index, section in enumerate(doc.sections, start=1):
        headers.append({"section": index, "text": "\n".join(p.text for p in section.header.paragraphs if p.text.strip())})
        footers.append({"section": index, "text": "\n".join(p.text for p in section.footer.paragraphs if p.text.strip())})

    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        media = [n for n in names if n.startswith("word/media/")]
        comments = "word/comments.xml" in names
        footnotes = "word/footnotes.xml" in names
        endnotes = "word/endnotes.xml" in names
        document_xml = archive.read("word/document.xml")
        tracked_changes = b"<w:ins" in document_xml or b"<w:del" in document_xml

    return {
        "path": str(path),
        "blocks": blocks,
        "headers": headers,
        "footers": footers,
        "media": media,
        "has_comments": comments,
        "has_footnotes": footnotes,
        "has_endnotes": endnotes,
        "has_tracked_changes": tracked_changes,
    }


if __name__ == "__main__":
    result = extract(Path(sys.argv[1]))
    print(json.dumps(result, ensure_ascii=False, indent=2))
